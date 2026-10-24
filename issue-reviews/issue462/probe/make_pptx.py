# /// script
# dependencies = ["python-pptx>=1.0", "python-docx>=1.1"]
# ///
"""Build throwaway readonly-probe.pptx/.docx whose media are the marked pictures, stored byte for byte.

Slide 1: marked PNG. Slide 2: SVG (asvg:svgBlip + marked PNG fallback, as Office stores SVG).
Slide 3: marked EMF. Slide 4: three pictures (PNG, PNG from marked2, SVG) for multi-selection.
"""
import copy
from pathlib import Path

from docx import Document
from docx.shared import Cm
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Cm as PCm

d = Path(__file__).parent
m, m2 = d / "marked", d / "marked2"
prs = Presentation()
blank = prs.slide_layouts[6]
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
svg_n = [0]


def add_svg(slide, left, svg_path, png_fallback):
    pic = slide.shapes.add_picture(str(png_fallback), left, PCm(3), width=PCm(10))
    svg_n[0] += 1
    part = Part(PackURI(f"/ppt/media/probe{svg_n[0]}.svg"), "image/svg+xml", prs.part.package, svg_path.read_bytes())
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.xpath(".//a:blip")[0]
    a = blip.nsmap.get("a") or "http://schemas.openxmlformats.org/drawingml/2006/main"
    ext_lst = etree.SubElement(blip, f"{{{a}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{a}}}ext", uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(ext, f"{{{ASVG}}}svgBlip", nsmap={"asvg": ASVG})
    svg_blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    return pic


s = prs.slides.add_slide(blank); s.shapes.add_picture(str(m / "marked.png"), PCm(2), PCm(3), width=PCm(12))
s = prs.slides.add_slide(blank); add_svg(s, PCm(2), m / "marked.svg", m / "marked.png")
s = prs.slides.add_slide(blank); s.shapes.add_picture(str(m / "marked.emf"), PCm(2), PCm(3), width=PCm(12))
s = prs.slides.add_slide(blank)
s.shapes.add_picture(str(m / "marked.png"), PCm(0.5), PCm(3), width=PCm(8))
s.shapes.add_picture(str(m2 / "marked.png"), PCm(9), PCm(3), width=PCm(8))
add_svg(s, PCm(17.5), m2 / "marked.svg", m2 / "marked.png").width = PCm(7)
prs.save(d / "readonly-probe.pptx")

doc = Document()
doc.add_picture(str(m / "marked.png"), width=Cm(10))
doc.save(d / "readonly-probe.docx")
print("ok")

# python-pptx files EMF as image2.wmf / image/x-wmf; store it as a real .emf part.
import zipfile
src = d / "readonly-probe.pptx"
data = {i.filename: zipfile.ZipFile(src).read(i) for i in zipfile.ZipFile(src).infolist()}
out = {}
for name, b in data.items():
    if name.endswith(".wmf"):
        name = name[:-4] + ".emf"
    elif name.endswith(".rels") or name == "[Content_Types].xml":
        b = b.replace(b"image2.wmf", b"image2.emf").replace(b'Extension="wmf" ContentType="image/x-wmf"', b'Extension="emf" ContentType="image/x-emf"')
    out[name] = b
with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as z:
    for name, b in out.items():
        z.writestr(name, b)
